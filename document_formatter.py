from io import BytesIO
from html import escape
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from fpdf import FPDF


def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
        "\u2013": "-", "\u2014": "-", "\u00a0": " "
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.strip()


def _lines(text):
    return [line.strip() for line in sanitize_text(text).splitlines()]


def format_html_preview(text: str) -> str:
    blocks = []
    for line in _lines(text):
        if not line:
            continue
        e = escape(line)
        if re.match(r"^\d+[\.\)]\s+", line):
            blocks.append(f"<h3>{e}</h3>")
        elif len(line) <= 100 and line.upper() == line and any(c.isalpha() for c in line):
            blocks.append(f"<h2>{e}</h2>")
        elif line.endswith(":") and len(line) <= 100:
            blocks.append(f"<h3>{e}</h3>")
        else:
            blocks.append(f"<p>{e}</p>")
    return "\n".join(blocks)


def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in _lines(text):
        if not line:
            continue
        if len(line) <= 100 and (line.upper() == line or re.match(r"^\d+[\.\)]\s+", line)):
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(6)

    terms = []
    for line in _lines(text):
        if ":" in line and len(line) < 180:
            key, value = line.split(":", 1)
            if key.strip() and value.strip():
                terms.append((key.strip(), value.strip()))

    if terms:
        doc.add_page_break()
        h = doc.add_paragraph()
        r = h.add_run("Key Terms")
        r.bold = True
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "Term"
        table.rows[0].cells[1].text = "Details"
        for key, value in terms[:20]:
            cells = table.add_row().cells
            cells[0].text = key
            cells[1].text = value

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-generated draft. Review before use.").italic = True

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, doc_type):
        super().__init__()
        self.doc_type = doc_type

    def header(self):
        self.set_font("Helvetica", "B", 13)
        self.cell(0, 8, self.doc_type.upper(), align="C")
        self.ln(8)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, "LegalEase - AI-generated draft. Review before use.", align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalEasePDF(doc_type)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)

    for line in _lines(text):
        if not line:
            pdf.ln(3)
            continue
        heading = len(line) <= 100 and (
            line.upper() == line or re.match(r"^\d+[\.\)]\s+", line)
        )
        if heading:
            pdf.set_font("Helvetica", "B", 11)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(pdf.epw, 7, line)
            pdf.set_font("Helvetica", size=11)
        else:
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(pdf.epw, 6, line)
            pdf.ln(1)
    return bytes(pdf.output())
